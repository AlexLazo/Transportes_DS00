"""Lectura tolerante de los Excel de la operación.

Los archivos cambiaron de formato durante el año (columnas nuevas, acentos, nombres de hoja), así que
todo se lee por NOMBRE de encabezado normalizado y no por posición fija.
"""
import re
import unicodedata
import warnings
from collections import Counter
from datetime import date, datetime

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


# ───────────── normalizadores ─────────────

def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def hnorm(v) -> str:
    """Encabezado/nombre normalizado: mayúsculas, sin acentos, espacios colapsados."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", sin_acentos(str(v)).upper()).strip()


def texto(v, maxlen=None):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s == "0":
        return None
    return s[:maxlen] if maxlen else s


def nombre_n(v):
    s = hnorm(v)
    return s or None


def ruta_n(v):
    if v is None:
        return None
    s = re.sub(r"\s+", "", str(v)).upper()[:40]
    return s if s and s != "0" else None


def a_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def a_int(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        n = int(float(str(v).replace(",", "").strip()))
    except (ValueError, TypeError):
        return None
    return n if n != 0 else None


def a_float(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return None


def si_no(v):
    s = hnorm(v)
    return s if s in ("SI", "NO") else None


def periodo_de(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def mapa_encabezados(fila, desde=0, hasta=None):
    """{encabezado_normalizado: índice}; si un nombre se repite se queda el primero."""
    out = {}
    for i in range(desde, len(fila) if hasta is None else min(hasta, len(fila))):
        h = hnorm(fila[i])
        if h and h not in out:
            out[h] = i
    return out


def celda(fila, mapa, *nombres):
    for n in nombres:
        i = mapa.get(n)
        if i is not None and i < len(fila):
            return fila[i]
    return None


def abrir(path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def hoja(wb, prefijo):
    """Primera hoja cuyo nombre normalizado empieza con `prefijo` (ignora acentos y mayúsculas)."""
    p = hnorm(prefijo)
    for n in wb.sheetnames:
        if hnorm(n).startswith(p):
            return wb[n]
    return None


# ───────────── archivo diario ─────────────

def parse_diario(path):
    """Hoja 'Asignacion' → tres bloques: Data Facturación (SAP), Validación de rutas y Cambios de cabecera."""
    wb = abrir(path)
    ws = hoja(wb, "asignacion")
    if ws is None:
        raise ValueError("No tiene hoja 'Asignacion'")
    filas = list(ws.iter_rows(min_row=1, max_row=3500, max_col=40, values_only=True))
    wb.close()

    fila_titulos = None
    for i, f in enumerate(filas[:8]):
        if any(hnorm(c).startswith("DATA FACTURACION") for c in f):
            fila_titulos = i
            break
    if fila_titulos is None:
        raise ValueError("No se encontró el bloque 'Data Facturacion'")
    titulos = {}
    for j, c in enumerate(filas[fila_titulos]):
        h = hnorm(c)
        if h.startswith("DATA FACTURACION"):
            titulos["fact"] = j
        elif h.startswith("VALIDACION DE RUTAS"):
            titulos["val"] = j
        elif h.startswith("CAMBIOS DE CABECERA"):
            titulos["cab"] = j
    inicios = sorted(titulos.values())

    def rango(clave):
        ini = titulos.get(clave)
        if ini is None:
            return None, None
        sig = [x for x in inicios if x > ini]
        return ini, (sig[0] if sig else 40)

    encabezado = filas[fila_titulos + 1]
    datos = filas[fila_titulos + 2:]
    avisos = []

    # Data Facturación (SAP)
    a, b = rango("fact")
    m = mapa_encabezados(encabezado, a, b)
    fact = []
    for f in datos:
        fecha = a_fecha(celda(f, m, "FECHA"))
        transp = a_int(celda(f, m, "TRANSPORTE"))
        if fecha is None or transp is None:
            continue
        fact.append(dict(
            fecha=fecha, transporte=transp, ruta=ruta_n(celda(f, m, "RUTA")),
            viaje=a_int(celda(f, m, "VIAJE")), st=a_int(celda(f, m, "ST")),
            cantidad=a_float(celda(f, m, "CANTIDAD")),
        ))
    if not fact:
        raise ValueError("El bloque 'Data Facturacion' no tiene filas")
    fecha_dia = Counter(r["fecha"] for r in fact).most_common(1)[0][0]
    if len({r["fecha"] for r in fact}) > 1:
        avisos.append("Data Facturación trae más de una fecha")

    # Validación de rutas
    val = []
    a, b = rango("val")
    if a is not None:
        m = mapa_encabezados(encabezado, a, b)
        vistos = set()
        for f in datos:
            r = ruta_n(celda(f, m, "RUTA"))
            if not r or r in vistos:
                continue
            vistos.add(r)
            tipo = texto(celda(f, m, "TIPO"), 30)
            if tipo and hnorm(tipo) == "FIAJA":
                tipo = "FIJA"
            val.append(dict(
                fecha=fecha_dia, ruta=r, estatus=texto(celda(f, m, "ESTATUS"), 200), tipo=tipo,
                contratista=texto(celda(f, m, "CONTRATISTA"), 120),
                supervisor=texto(celda(f, m, "SUPERVISOR"), 120),
            ))

    # Cambios de cabecera
    cab = []
    a, b = rango("cab")
    if a is not None:
        m = mapa_encabezados(encabezado, a, b)
        vistos = set()
        for f in datos:
            t = a_int(celda(f, m, "TRANSPORTE"))
            if t is None or t in vistos:
                continue
            vistos.add(t)
            cab.append(dict(
                fecha=fecha_dia, transporte=t,
                ruta_original=ruta_n(celda(f, m, "RUTA")),
                cod_transportista=texto(celda(f, m, "TRANSPORTISTA"), 20),
                nueva_ruta=ruta_n(celda(f, m, "NUEVA RUTA")),
                conductor_sap=texto(celda(f, m, "CONDUCTOR SAP"), 20),
                tractor=texto(celda(f, m, "TRACTOR"), 20),
                viaje=a_int(celda(f, m, "VIAJE")),
                cajas=a_float(celda(f, m, "CAJAS")),
                contratista=texto(celda(f, m, "CONTRATISTA"), 120),
                comentario=texto(celda(f, m, "COMENTARIO")),
            ))
    return dict(fecha=fecha_dia, fact=fact, val=val, cab=cab, avisos=avisos)


# ───────────── excepciones (pago de dummies) ─────────────

def parse_excepciones(path):
    wb = abrir(path)
    ws = hoja(wb, "ddbb")
    if ws is None:
        wb.close()
        raise ValueError("No tiene hoja 'DDBB'")
    filas = list(ws.iter_rows(values_only=True, max_col=24))
    wb.close()
    m = mapa_encabezados(filas[0])
    rows = []
    for n, f in enumerate(filas[1:], start=2):
        fecha = a_fecha(celda(f, m, "FECHA"))
        if fecha is None:
            continue
        contratista = texto(celda(f, m, "CONTRATISTA"), 120)
        rows.append(dict(
            fila=n, fecha=fecha, periodo=periodo_de(fecha),
            transporte=a_int(celda(f, m, "TRANSPORTE")),
            ruta=ruta_n(celda(f, m, "RUTA")),
            ruta_liquidada=ruta_n(celda(f, m, "RUTA LIQUIDADA")),
            dummy=si_no(celda(f, m, "DUMMY?")), recarga=si_no(celda(f, m, "RECARGA?")),
            excepcion=si_no(celda(f, m, "EXCEPCION")), ruta_fija=si_no(celda(f, m, "RUTA FIJA")),
            spot1=si_no(celda(f, m, "SPOT 1")), spot_inc=si_no(celda(f, m, "SPOT INCREMENTAL")),
            envases=si_no(celda(f, m, "ENVASES")),
            contratista=contratista, contratista_n=nombre_n(contratista),
            cod_od=texto(celda(f, m, "CODIGO OD"), 20),
            firmada=texto(celda(f, m, "FIRMADA"), 40), cod=texto(celda(f, m, "COD"), 40),
            comentario=texto(celda(f, m, "COMENTARIO")),
            quincena=a_int(celda(f, m, "QUINCENA")),
        ))
    if not rows:
        raise ValueError("DDBB sin filas con fecha")
    # El mes del archivo es el más frecuente (hay filas rezagadas del mes anterior).
    periodo = Counter(r["periodo"] for r in rows).most_common(1)[0][0]
    return dict(rows=rows, periodo=periodo, max_fecha=max(r["fecha"] for r in rows))


# ───────────── ausentismos ─────────────

def parse_ausentismos(path):
    wb = abrir(path)
    ws = hoja(wb, "registros")
    if ws is None:
        wb.close()
        raise ValueError("No tiene hoja 'REGISTROS'")
    filas = list(ws.iter_rows(values_only=True, max_col=16))
    wb.close()
    m = mapa_encabezados(filas[0])
    rows = []
    for n, f in enumerate(filas[1:], start=2):
        fecha = a_fecha(celda(f, m, "FECHA"))
        if fecha is None:
            continue
        contratista = texto(celda(f, m, "CONTRATISTA"), 120)
        rows.append(dict(
            fila=n, fecha=fecha, periodo=periodo_de(fecha),
            ruta=ruta_n(celda(f, m, "RUTA ASIGNADA")),
            contratista=contratista, contratista_n=nombre_n(contratista),
            motivo=texto(celda(f, m, "MOTIVO"), 120),
            tripulacion=a_float(celda(f, m, "TRIPULACION")),
            costo_ruta=a_float(celda(f, m, "COSTO DE RUTA")),
            dias_trab=a_float(celda(f, m, "DIAS TRABAJADOS")),
            costo_dia=a_float(celda(f, m, "COSTO X DIA")),
            comentario=texto(celda(f, m, "COMENTARIO")),
        ))
    if not rows:
        raise ValueError("REGISTROS sin filas con fecha")
    periodo = Counter(r["periodo"] for r in rows).most_common(1)[0][0]
    return dict(rows=rows, periodo=periodo, max_fecha=max(r["fecha"] for r in rows))


# ───────────── rutas fijas del mes ─────────────

def _filas_rutas(ws, fuente):
    filas = list(ws.iter_rows(values_only=True, max_col=24))
    if not filas:
        return None
    m = mapa_encabezados(filas[0])
    if "CODIGO RUTA" not in m:
        return None
    out = []
    for f in filas[1:]:
        ruta = ruta_n(celda(f, m, "CODIGO RUTA"))
        if not ruta:
            continue
        anio, mes = a_int(celda(f, m, "ANO")), a_int(celda(f, m, "MES DIMENSIONADO"))
        out.append(dict(
            anio=anio, mes=mes, ruta=ruta,
            canal=texto(celda(f, m, "CANAL"), 60),
            cod_sap=texto(celda(f, m, "CODIGO SAP PROVEEDOR"), 20),
            personas=a_int(celda(f, m, "CANTIDAD DE PERSONAS EN RUTA")),
            categoria=texto(celda(f, m, "CATEGORIA"), 10),
            estatus=texto(celda(f, m, "ESTATUS"), 30), tipo=texto(celda(f, m, "TIPO"), 30),
            transportista=texto(celda(f, m, "TRANSPORTISTA"), 120),
            supervisor=texto(celda(f, m, "SUP"), 120),
            vehiculo=texto(celda(f, m, "SV"), 20), placa=texto(celda(f, m, "PLACA"), 20),
            fuente=fuente,
        ))
    return out


def parse_rutas_fijas(path):
    """Devuelve {periodo: [filas]} (la hoja visible 'RUTAS X SUP' y, si existen, hojas ocultas de otro mes
    como 'Rutas Fijas JULIO') y la lista de rutas desactivadas."""
    wb = abrir(path)
    por_periodo = {}
    for ws in wb.worksheets:
        h = hnorm(ws.title)
        if h == "RUTAS X SUP":
            fuente = "VISIBLE"
        elif h.startswith("RUTAS FIJAS"):
            fuente = "OCULTA"
        else:
            continue
        filas = _filas_rutas(ws, fuente)
        if not filas:
            continue
        mode = Counter((r["anio"], r["mes"]) for r in filas if r["anio"] and r["mes"]).most_common(1)
        if not mode:
            continue
        (anio, mes), _ = mode[0]
        periodo = f"{anio:04d}-{mes:02d}"
        por_periodo.setdefault(periodo, []).extend(r for r in filas if (r["anio"], r["mes"]) == (anio, mes))
    desact = []
    ws = hoja(wb, "rutas desactivadas")
    if ws is not None:
        filas = list(ws.iter_rows(values_only=True, max_col=7))
        for f in filas[1:]:
            r = ruta_n(f[0])
            if r:
                desact.append(dict(ruta=r, estatus=texto(f[1], 30), tripulacion=a_int(f[2]),
                                   comentario=texto(f[3])))
    wb.close()
    if not por_periodo:
        raise ValueError("No se encontró 'RUTAS X SUP' con datos")
    return dict(por_periodo=por_periodo, desactivadas=desact)


# ───────────── control de AJ de la supervisora ─────────────

def _flag(v):
    """SI/NO tolerante: 'N0' (cero) se corrige a NO; devuelve (valor, aviso|None)."""
    h = hnorm(v)
    if h in ("SI", "NO"):
        return h, None
    c = h.replace("0", "O")
    if c in ("SI", "NO"):
        return c, f"«{v}» corregido a {c}"
    return None, (f"valor no válido «{v}»" if h else None)


def _filas_desde(ws, clave="FECHA"):
    """(mapa de encabezados, [(n_fila, fila)]) a partir de la fila que contiene el encabezado `clave`."""
    filas = list(ws.iter_rows(values_only=True, max_col=20))
    for i, f in enumerate(filas[:8]):
        m = mapa_encabezados(f)
        if clave in m:
            return m, list(enumerate(filas[i + 1:], start=i + 2))
    return None, []


def parse_cargas_aj(path):
    wb = abrir(path)
    cargas, recargas, recolec, vistos = [], [], [], set()
    for ws in wb.worksheets:
        h = hnorm(ws.title)
        if h.startswith("CARGAS"):
            m, filas = _filas_desde(ws)
            for n, f in filas:
                fecha, t = a_fecha(celda(f, m, "FECHA")), a_int(celda(f, m, "TRANSPORTE"))
                if fecha is None or t is None or t in vistos:
                    continue
                vistos.add(t)
                avisos, vals = [], {}
                for col, clave in (("CARGA/DIA", "carga_dia"), ("DUMMY", "dummy"), ("RECARGA", "recarga"),
                                   ("REPOSICION", "reposicion"), ("RECOLECCION", "recoleccion")):
                    vals[clave], av = _flag(celda(f, m, col))
                    if av:
                        avisos.append(f"{col}: {av}")
                cargas.append(dict(
                    fila=n, fecha=fecha, periodo=periodo_de(fecha), transporte=t,
                    ruta=ruta_n(celda(f, m, "RUTA")), viaje=a_int(celda(f, m, "VIAJE")),
                    cantidad=a_float(celda(f, m, "CANTIDAD")),
                    transportista=texto(celda(f, m, "TRANSPORTISTA"), 20),
                    aviso="; ".join(avisos)[:200] or None, **vals))
        elif h.startswith("RECARGAS"):
            m, filas = _filas_desde(ws)
            for n, f in filas:
                fecha = a_fecha(celda(f, m, "FECHA"))
                if fecha is None:
                    continue
                recargas.append(dict(fila=n, fecha=fecha, periodo=periodo_de(fecha),
                                     transporte=a_int(celda(f, m, "TRANSPORTE")), ruta=ruta_n(celda(f, m, "RUTA")),
                                     cantidad=a_float(celda(f, m, "CANTIDAD")),
                                     cliente=texto(celda(f, m, "CLIENTE"), 120)))
        elif h.startswith("RECOLEC"):
            m, filas = _filas_desde(ws)
            for n, f in filas:
                fecha = a_fecha(celda(f, m, "FECHA"))
                if fecha is None:
                    continue
                recolec.append(dict(fila=n, fecha=fecha, periodo=periodo_de(fecha),
                                    ruta=ruta_n(celda(f, m, "RUTA")),
                                    personal=texto(celda(f, m, "RECOLECCION DE ENVASE/RACKS"), 120),
                                    cliente=texto(celda(f, m, "CLIENTE"), 120),
                                    envases=a_float(celda(f, m, "ENVASES")), racks=a_float(celda(f, m, "RACKS"))))
    wb.close()
    if not cargas:
        raise ValueError("No se encontró ninguna hoja 'Cargas ...' con datos")
    fechas = [r["fecha"] for r in cargas]
    return dict(cargas=cargas, recargas=recargas, recolec=recolec,
                periodos=sorted({r["periodo"] for r in cargas}), max_fecha=max(fechas), min_fecha=min(fechas))
